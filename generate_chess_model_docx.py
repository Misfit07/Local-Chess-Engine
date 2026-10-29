import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor

def build_model_docx():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run("Java Chess Engine — chess.model Package")
    title_run.font.name = "Segoe UI"
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(27, 54, 93)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(8)
    sub_run = sub_p.add_run("Complete Source Code: Board Representation, Move Representation, Pieces & Game State")
    sub_run.font.name = "Segoe UI"
    sub_run.font.size = Pt(9)
    sub_run.font.color.rgb = RGBColor(90, 100, 115)

    files = [
        ("1. PieceColor.java", "src/chess/model/PieceColor.java",
         "Color enum defining WHITE and BLACK with opposite() color toggling and pawn rank/direction helpers."),
        
        ("2. PieceType.java", "src/chess/model/PieceType.java",
         "Piece type enum (PAWN, KNIGHT, BISHOP, ROOK, QUEEN, KING) with character symbols and base material values."),
        
        ("3. Square.java", "src/chess/model/Square.java",
         "0-63 square coordinate helpers, rank/file math, algebraic name conversion (e.g. 28 -> 'e4'), and bounds checks."),
        
        ("4. MoveFlag.java", "src/chess/model/MoveFlag.java",
         "Move categorization flag enum (NORMAL, CASTLE, EN_PASSANT, PROMOTION) driving special move mechanics."),
        
        ("5. Move.java", "src/chess/model/Move.java",
         "Immutable move record storing from/to squares, moving piece, captured piece, promotion piece, flag, and UCI string."),
        
        ("6. Piece.java", "src/chess/model/Piece.java",
         "Abstract piece base class providing flyweight piece caching and polymorphic move generation contract."),
        
        ("7. Pawn.java", "src/chess/model/Pawn.java",
         "Pawn move generation handling single push, initial double push, diagonal captures, en passant, and 4-way promotion."),
        
        ("8. Knight.java", "src/chess/model/Knight.java",
         "Knight move generation using 8 discrete coordinate offsets with board-edge wrap prevention."),
        
        ("9. Bishop.java", "src/chess/model/Bishop.java",
         "Bishop move generation using diagonal ray marching along 4 diagonal offsets (+9, +7, -7, -9)."),
        
        ("10. Rook.java", "src/chess/model/Rook.java",
         "Rook move generation using orthogonal ray marching along 4 orthogonal offsets (+8, -8, +1, -1)."),
        
        ("11. Queen.java", "src/chess/model/Queen.java",
         "Queen move generation combining all 8 orthogonal and diagonal sliding ray directions."),
        
        ("12. King.java", "src/chess/model/King.java",
         "King move generation with 8 one-step offsets and castling legality checks (path clearance and check safety)."),
        
        ("13. Board.java", "src/chess/model/Board.java",
         "Core 8x8 chessboard representation, position state, reversible makeMove/unmakeMove with undo stack, and incremental Zobrist XOR hashing."),
        
        ("14. GameState.java", "src/chess/model/GameState.java",
         "Game termination status enum (ACTIVE, CHECKMATE, STALEMATE, DRAW_FIFTY_MOVES, DRAW_INSUFFICIENT_MATERIAL, DRAW_THREEFOLD_REPETITION)."),
        
        ("15. GameStateDetector.java", "src/chess/model/GameStateDetector.java",
         "Comprehensive game-ending evaluation verifying checkmate, stalemate, fifty-move rule, threefold repetition, and insufficient material.")
    ]

    base_dir = r"c:\Users\jeffe\Downloads\javaChessEngine"

    for display_title, rel_path, desc in files:
        full_path = os.path.join(base_dir, rel_path)
        if not os.path.exists(full_path):
            continue

        with open(full_path, "r", encoding="utf-8") as f:
            code_content = f.read()

        h2 = doc.add_paragraph()
        h2.paragraph_format.space_before = Pt(8)
        h2.paragraph_format.space_after = Pt(1)
        h2.paragraph_format.keep_with_next = True
        h2_run = h2.add_run(f"═══ {display_title} ═══")
        h2_run.font.name = "Segoe UI"
        h2_run.font.size = Pt(10.5)
        h2_run.font.bold = True
        h2_run.font.color.rgb = RGBColor(43, 87, 154)

        desc_p = doc.add_paragraph()
        desc_p.paragraph_format.space_before = Pt(0)
        desc_p.paragraph_format.space_after = Pt(2)
        desc_p.paragraph_format.keep_with_next = True
        desc_run = desc_p.add_run("// " + desc)
        desc_run.font.name = "Consolas"
        desc_run.font.size = Pt(7.5)
        desc_run.font.italic = True
        desc_run.font.color.rgb = RGBColor(100, 100, 100)

        lines = code_content.splitlines()
        prev_blank = False
        for line in lines:
            is_blank = not line.strip()
            if is_blank and prev_blank:
                continue
            prev_blank = is_blank

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(8.5)
            
            run = p.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(7.5)
            run.font.color.rgb = RGBColor(25, 25, 25)

    out_path = os.path.join(base_dir, "JavaChessEngine_ChessModel_Code.docx")
    doc.save(out_path)
    print(f"Successfully generated: {out_path}")

if __name__ == "__main__":
    build_model_docx()
