import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor

def build_docx():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(1)
    title_run = title_p.add_run("Java Chess Engine — Source Code & Architecture")
    title_run.font.name = "Segoe UI"
    title_run.font.size = Pt(12)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 60, 120)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(4)
    sub_run = sub_p.add_run("Target: JDK 17 | Plain Eclipse-Compatible Project | Zero External Dependencies")
    sub_run.font.name = "Segoe UI"
    sub_run.font.size = Pt(7.5)
    sub_run.font.color.rgb = RGBColor(100, 100, 100)

    packages = [
        ("Project Configuration", [
            (".project", ".project", "Eclipse project definition with Java builder and nature."),
            (".classpath", ".classpath", "Eclipse classpath configuring src/ directory, JavaSE-17 JRE, and JUnit 5 container.")
        ]),
        ("Package: chess.model", [
            ("PieceColor.java", "src/chess/model/PieceColor.java", "Color enum with opposite color and pawn rank/direction helpers."),
            ("PieceType.java", "src/chess/model/PieceType.java", "Piece type enum with character symbols and base material values."),
            ("Square.java", "src/chess/model/Square.java", "0-63 square coordinate helpers, rank/file math, algebraic name conversion."),
            ("MoveFlag.java", "src/chess/model/MoveFlag.java", "Move flag enum (NORMAL, CASTLE, EN_PASSANT, PROMOTION)."),
            ("Move.java", "src/chess/model/Move.java", "Immutable move record with from, to, pieces, flag, and UCI representation."),
            ("Piece.java", "src/chess/model/Piece.java", "Abstract piece base class with flyweight caching and polymorphic move generation."),
            ("Pawn.java", "src/chess/model/Pawn.java", "Pawn moves: single/double push, diagonal captures, en passant, promotions."),
            ("Knight.java", "src/chess/model/Knight.java", "Knight move generation using 8 coordinate offsets."),
            ("Bishop.java", "src/chess/model/Bishop.java", "Bishop move generation using diagonal sliding rays."),
            ("Rook.java", "src/chess/model/Rook.java", "Rook move generation using orthogonal sliding rays."),
            ("Queen.java", "src/chess/model/Queen.java", "Queen move generation combining 8 orthogonal and diagonal sliding rays."),
            ("King.java", "src/chess/model/King.java", "King moves with 8 offsets and castling legality/path checks."),
            ("Board.java", "src/chess/model/Board.java", "8x8 array representation, position state, undo history, makeMove/unmakeMove mirror, and incremental XOR Zobrist."),
            ("GameState.java", "src/chess/model/GameState.java", "Game termination status enum (ACTIVE, CHECKMATE, STALEMATE, 50-move, insufficient material, repetition)."),
            ("GameStateDetector.java", "src/chess/model/GameStateDetector.java", "Rule evaluation for checkmate, stalemate, fifty-move rule, insufficient material, and threefold repetition.")
        ]),
        ("Package: chess.movegen", [
            ("AttackDetector.java", "src/chess/movegen/AttackDetector.java", "Attack detection for pawns, knights, kings, and sliding pieces."),
            ("MoveGenerator.java", "src/chess/movegen/MoveGenerator.java", "Pseudo-legal move generation and check filtering for legal moves.")
        ]),
        ("Package: chess.search", [
            ("Evaluator.java", "src/chess/search/Evaluator.java", "Evaluation interface defining relative and absolute evaluations."),
            ("MaterialPstEvaluator.java", "src/chess/search/MaterialPstEvaluator.java", "Material balance + PeSTO piece-square tables with mirrored black perspective."),
            ("MoveOrderer.java", "src/chess/search/MoveOrderer.java", "MVV-LVA move ordering heuristic sorting captures and promotions first."),
            ("RootMoveScore.java", "src/chess/search/RootMoveScore.java", "Data record storing root moves and scores for top-N ranking."),
            ("SearchResult.java", "src/chess/search/SearchResult.java", "Container for best move, evaluation, depth, nodes, elapsed time, and top moves."),
            ("Search.java", "src/chess/search/Search.java", "Negamax alpha-beta search, quiescence, iterative deepening, time control, TT cutoffs, and mate distance adjustment.")
        ]),
        ("Package: chess.hash", [
            ("Zobrist.java", "src/chess/hash/Zobrist.java", "64-bit random tables for pieces x squares, side, 4 castling flags, 8 en passant files."),
            ("RepetitionTracker.java", "src/chess/hash/RepetitionTracker.java", "Two-tier repetition tracker for real-game history and search-time stack."),
            ("TranspositionTable.java", "src/chess/hash/TranspositionTable.java", "Transposition table storing depths, scores, exact/lower/upper flags, and PV moves keyed by Zobrist hash.")
        ]),
        ("Package: chess.ui", [
            ("BoardPanel.java", "src/chess/ui/BoardPanel.java", "Swing 8x8 chessboard with wood tones, Unicode pieces, highlights, click-to-move, and background AI worker."),
            ("ChessFrame.java", "src/chess/ui/ChessFrame.java", "Main GUI window with menus (Game, Mode, Difficulty, FEN) and live status bar."),
            ("ConsoleGame.java", "src/chess/ui/ConsoleGame.java", "Interactive CLI game with ASCII board rendering, move parsing, and command loop.")
        ])
    ]

    base_dir = r"c:\Users\jeffe\Downloads\javaChessEngine"

    for pkg_title, files in packages:
        h1 = doc.add_paragraph()
        h1.paragraph_format.space_before = Pt(4)
        h1.paragraph_format.space_after = Pt(1)
        h1_run = h1.add_run("═══ " + pkg_title + " ═══")
        h1_run.font.name = "Segoe UI"
        h1_run.font.size = Pt(9.5)
        h1_run.font.bold = True
        h1_run.font.color.rgb = RGBColor(43, 87, 154)

        for name, rel_path, desc in files:
            full_path = os.path.join(base_dir, rel_path)
            if not os.path.exists(full_path):
                continue
            
            with open(full_path, "r", encoding="utf-8") as f:
                code_content = f.read()

            h2 = doc.add_paragraph()
            h2.paragraph_format.space_before = Pt(3)
            h2.paragraph_format.space_after = Pt(1)
            h2_run = h2.add_run(f"── {name} ({rel_path}) ──")
            h2_run.font.name = "Segoe UI"
            h2_run.font.size = Pt(8)
            h2_run.font.bold = True
            h2_run.font.color.rgb = RGBColor(40, 40, 40)

            desc_p = doc.add_paragraph()
            desc_p.paragraph_format.space_before = Pt(0)
            desc_p.paragraph_format.space_after = Pt(1)
            desc_run = desc_p.add_run("// " + desc)
            desc_run.font.name = "Consolas"
            desc_run.font.size = Pt(6)
            desc_run.font.italic = True
            desc_run.font.color.rgb = RGBColor(110, 110, 110)

            lines = code_content.splitlines()
            prev_blank = False
            for line in lines:
                is_blank = not line.strip()
                if is_blank and prev_blank:
                    continue
                prev_blank = is_blank

                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(7.0)
                
                run = p.add_run(line if line else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(6.2)
                run.font.color.rgb = RGBColor(25, 25, 25)

    out_path = os.path.join(base_dir, "JavaChessEngine_Complete_Codebase.docx")
    doc.save(out_path)
    print(f"Successfully generated: {out_path}")

if __name__ == "__main__":
    build_docx()
