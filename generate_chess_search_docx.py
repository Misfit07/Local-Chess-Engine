import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor

def build_search_docx():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    title_run = title_p.add_run("Java Chess Engine — chess.search Package")
    title_run.font.name = "Segoe UI"
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(27, 54, 93)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(8)
    sub_run = sub_p.add_run("Complete Source Code: Search Engine, Minimax / Alpha-Beta Pruning, Evaluators & Move Ordering")
    sub_run.font.name = "Segoe UI"
    sub_run.font.size = Pt(9)
    sub_run.font.color.rgb = RGBColor(90, 100, 115)

    files = [
        ("1. Evaluator.java", "src/chess/search/Evaluator.java",
         "Public evaluation interface defining relative (side to move) and absolute (White perspective) evaluation methods."),
        
        ("2. MaterialPstEvaluator.java", "src/chess/search/MaterialPstEvaluator.java",
         "Static evaluation implementation combining base material balance and PeSTO Piece-Square Tables (PST) with mirrored black perspective."),
        
        ("3. MoveOrderer.java", "src/chess/search/MoveOrderer.java",
         "Move ordering heuristic using Most Valuable Victim - Least Valuable Aggressor (MVV-LVA) to test promising captures and promotions first for maximum alpha-beta pruning."),
        
        ("4. RootMoveScore.java", "src/chess/search/RootMoveScore.java",
         "Data transfer record storing root-level moves alongside their minimax evaluation scores for principal variation analysis and multi-PV ranking."),
        
        ("5. SearchResult.java", "src/chess/search/SearchResult.java",
         "Immutable result container holding best move, evaluation score, search depth reached, total nodes searched, elapsed time, and top candidate moves."),
        
        ("6. Search.java", "src/chess/search/Search.java",
         "Core search engine implementing Negamax with Alpha-Beta pruning, Quiescence search, Iterative Deepening, Transposition Table cutoffs, Time Management, and Mate Distance scoring.")
    ]

    base_dir = r"c:\Users\jeffe\Downloads\javaChessEngine"

    for display_title, rel_path, desc in files:
        full_path = os.path.join(base_dir, rel_path)
        if not os.path.exists(full_path):
            continue

        with open(full_path, "r", encoding="utf-8") as f:
            code_content = f.read()

        h2 = doc.add_paragraph()
        h2.paragraph_format.space_before = Pt(8)
        h2.paragraph_format.space_after = Pt(1)
        h2.paragraph_format.keep_with_next = True
        h2_run = h2.add_run(f"═══ {display_title} ═══")
        h2_run.font.name = "Segoe UI"
        h2_run.font.size = Pt(10.5)
        h2_run.font.bold = True
        h2_run.font.color.rgb = RGBColor(43, 87, 154)

        desc_p = doc.add_paragraph()
        desc_p.paragraph_format.space_before = Pt(0)
        desc_p.paragraph_format.space_after = Pt(2)
        desc_p.paragraph_format.keep_with_next = True
        desc_run = desc_p.add_run("// " + desc)
        desc_run.font.name = "Consolas"
        desc_run.font.size = Pt(7.5)
        desc_run.font.italic = True
        desc_run.font.color.rgb = RGBColor(100, 100, 100)

        lines = code_content.splitlines()
        prev_blank = False
        for line in lines:
            is_blank = not line.strip()
            if is_blank and prev_blank:
                continue
            prev_blank = is_blank

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(8.5)
            
            run = p.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(7.5)
            run.font.color.rgb = RGBColor(25, 25, 25)

    out_path = os.path.join(base_dir, "JavaChessEngine_ChessSearch_Code.docx")
    doc.save(out_path)
    print(f"Successfully generated: {out_path}")

if __name__ == "__main__":
    build_search_docx()
